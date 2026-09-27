"""Verification for how viewer.xlsx is BUILT (vpa/viewer_xlsx.generate).

    python -m tests.check_workbook      (from vehicle-price-analysis/)

check_viewer proves the workbook's formulas answer correctly in real Excel.
This suite proves the build around them is safe and honest, without Excel:

  * reused forecasts are the forecasts a fresh computation gives: after a
    price changes, a build that reuses the cache is cell-for-cell identical
    to a build from scratch, and only the touched model was recomputed;
  * a cache from other code, or a damaged one, is ignored rather than
    trusted;
  * progress is reported, only moves forward and ends at 100%;
  * a workbook held open (as Excel holds it) fails in a moment, before any
    work, and the file is untouched;
  * a save that dies part-way leaves the previous workbook exactly as it
    was, and no half-written file behind; a good save keeps the replaced
    file as viewer.previous.xlsx;
  * the data stamp matches the database it was built from and stops
    matching when the data changes - which is how the app knows it is stale;
  * the MyView sheet is never written to by a rebuild.

Everything runs on throwaway files under tests/synthetic/. ASCII output;
exit 0 only if all checks pass.
"""
from __future__ import annotations

import ctypes
import json
import shutil
import sqlite3
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

import vpa.viewer_xlsx as vx
from vpa.schema import make_record
from vpa.store import Store

ROOT = Path(__file__).resolve().parent
WORK = ROOT / "synthetic" / "workbook_work"
DB = WORK / "wb_test.db"

RESULTS: list[tuple[bool, str]] = []


def check(name: str, ok: bool, detail: str) -> None:
    RESULTS.append((bool(ok), name))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}: {detail}")


def seed(db: Path) -> None:
    """Three model lines, a trim each with plenty of prices, five monthly
    captures - enough for real forecasts rather than 'insufficient'."""
    raw = []
    lines = (("Ford", "Falcon", "XR6", 24000), ("Holden", "Commodore", "SS", 30000),
             ("Toyota", "Hilux", "SR5", 38000))
    for make, model, variant, base in lines:
        for i in range(10):
            for m, day in enumerate(("2026-01-10", "2026-02-10", "2026-03-10",
                                     "2026-04-10", "2026-05-10")):
                raw.append({"listing_id": f"{make[:2].upper()}-{i}", "make": make,
                            "model": model, "variant": variant, "year": 2014,
                            "price": base + i * 900 - m * 250,
                            "odometer": 90000 + i * 7000, "state": "NSW",
                            "seller_type": "dealer", "capture_date": day})
    recs = [r for r in (make_record(d, source="carsales") for d in raw) if r]
    with Store(db) as store:
        store.ingest(recs, source_file="(workbook-test)", adapter="test")


def cells(path: Path) -> dict:
    wb = load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or (ws.title, c.coordinate) == ("_Forecasts", "F1"):
                    continue                     # F1 = "generated <time>"
                v = c.value
                out[(ws.title, c.coordinate)] = (getattr(v, "text", v), c.number_format)
    return out


def build(db: Path, out: Path) -> tuple[Path, list]:
    log: list = []
    vx.generate(db, out, progress=lambda f, t: log.append((f, t)))
    return out, log


def recomputed(log: list) -> int:
    """How many forecasts a build actually computed (from its progress)."""
    n = [t for _, t in log if t.startswith("Forecasts ") and " of " in t]
    return len(n)


def hold_exclusively(path: Path):
    """Open the file the way Excel does - no sharing - and return a closer."""
    k32 = ctypes.windll.kernel32
    k32.CreateFileW.restype = ctypes.c_void_p
    h = k32.CreateFileW(str(path), 0x80000000, 0, None, 3, 0x80, None)
    if h in (None, ctypes.c_void_p(-1).value):
        raise OSError("could not hold the file")
    return lambda: k32.CloseHandle(ctypes.c_void_p(h))


def main() -> int:
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)
    seed(DB)

    print("\n== reused forecasts are the forecasts a fresh build gives ==")
    a = WORK / "a" / "viewer.xlsx"
    a.parent.mkdir()
    _, cold = build(DB, a)
    cache = json.loads((a.parent / "viewer_cache.json").read_text())
    n_fc = len(cache["rows"])
    check("a first build computes every forecast and caches it",
          recomputed(cold) == n_fc and n_fc >= 6,
          f"{recomputed(cold)} computed, {n_fc} cached")
    first = cells(a)
    _, warm = build(DB, a)
    check("an unchanged database rebuilds with no forecast recomputed",
          recomputed(warm) == 0, f"{recomputed(warm)} recomputed")
    check("...and the workbook is cell-for-cell what the first build wrote",
          cells(a) == first, f"{len(first):,} cells compared")

    with sqlite3.connect(DB) as conn:           # one Holden price moves
        conn.execute("UPDATE observations SET price_aud = price_aud + 1500 WHERE id = "
                     "(SELECT o.id FROM observations o JOIN vehicles v ON v.key=o.key "
                     " WHERE v.make='Holden' ORDER BY o.id LIMIT 1)")
    _, after = build(DB, a)
    b = WORK / "b" / "viewer.xlsx"
    b.parent.mkdir()
    build(DB, b)                                 # no cache at all
    check("after a price changes, only that model's forecasts are recomputed",
          recomputed(after) == 2,
          f"{recomputed(after)} recomputed (the Commodore line and its SS trim)")
    ca, cb = cells(a), cells(b)
    diff = [k for k in set(ca) | set(cb) if ca.get(k) != cb.get(k)]
    check("...and the cached build equals a build from scratch, every cell",
          not diff, f"{len(ca):,} cells, {len(diff)} differ {sorted(diff)[:3]}")
    fc_changed = [k for k in ca if k[0] == "_Forecasts" and first.get(k) != ca[k]]
    check("...which really did change the forecast (the test has teeth)",
          bool(fc_changed), f"{len(fc_changed)} _Forecasts cells moved")

    print("\n== a cache is trusted only when it is certainly right ==")
    cfile = a.parent / "viewer_cache.json"
    d = json.loads(cfile.read_text())
    d["sig"] = "other-code"
    cfile.write_text(json.dumps(d))
    _, other = build(DB, a)
    check("a cache written by different code is ignored",
          recomputed(other) == n_fc, f"{recomputed(other)} of {n_fc} recomputed")
    cfile.write_text("{not json")
    _, broken = build(DB, a)
    check("a damaged cache is ignored and the build still succeeds",
          recomputed(broken) == n_fc and cells(a) == cb,
          f"{recomputed(broken)} recomputed, output identical to a fresh build")

    print("\n== progress ==")
    fr = [f for f, _ in cold]
    check("progress only moves forward and ends at 100%",
          fr == sorted(fr) and fr[-1] == 1.0 and fr[0] == 0.0,
          f"{len(fr)} reports, {fr[0]} -> {fr[-1]}")
    stages = []
    for _, t in cold:
        s = t.split(" (")[0].split(" 1")[0].split(" of ")[0]
        s = "Forecasts" if s.startswith("Forecasts ") else s
        if not stages or stages[-1] != s:
            stages.append(s)
    check("...naming each stage in turn",
          {"Reading the database", "Forecasts", "Opening the current workbook",
           "Writing the data sheets", "Writing the views",
           "Saving viewer.xlsx", "Saved"} <= set(stages),
          " > ".join(stages))

    print("\n== a workbook open in Excel ==")
    before = a.read_bytes()
    release = hold_exclusively(a)
    log: list = []
    t0 = time.time()
    try:
        vx.generate(DB, a, progress=lambda f, t: log.append(t))
        refused = False
    except PermissionError:
        refused = True
    finally:
        release()
    took = time.time() - t0
    check("is refused at once, before any work",
          refused and took < 2 and not any(t.startswith("Forecasts") for t in log),
          f"PermissionError={refused} after {took:.2f}s, stages reached: {log}")
    check("...and the file is untouched", a.read_bytes() == before,
          f"{len(before):,} bytes identical")

    print("\n== a save that dies part-way ==")
    real = vx._write_xlsx

    def dies(wb, path, prog=None):
        Path(path).write_bytes(b"PK half a file")
        raise OSError("disk full (simulated)")
    vx._write_xlsx = dies
    try:
        try:
            vx.generate(DB, a)
            raised = False
        except OSError:
            raised = True
    finally:
        vx._write_xlsx = real
    leftovers = [p.name for p in a.parent.iterdir() if ".building" in p.name]
    check("leaves the previous workbook exactly as it was",
          raised and a.read_bytes() == before,
          f"error raised={raised}, viewer.xlsx {len(a.read_bytes()):,} bytes, unchanged")
    check("...and no half-written file behind", not leftovers, f"leftovers: {leftovers}")
    prev = a.with_name("viewer.previous.xlsx")
    build(DB, a)
    check("a good save keeps the replaced file as viewer.previous.xlsx",
          prev.exists() and prev.read_bytes() == before,
          f"{prev.name}: {prev.stat().st_size:,} bytes = the file it replaced")

    print("\n== the data stamp ==")
    st, fp = vx.workbook_stamp(a), vx.data_fingerprint(DB)
    check("the workbook is stamped with the data it was built from",
          st and st["data"] == fp, f"stamp {st and st['data']} vs database {fp}")
    with sqlite3.connect(DB) as conn:           # a change no total can see
        conn.execute("UPDATE vehicles SET state = 'VIC' WHERE rowid = 1")
    check("...and stops matching when the data changes",
          vx.workbook_stamp(a)["data"] != vx.data_fingerprint(DB),
          "one vehicle's state changed; stamp and database now differ")
    check("no file, or not a workbook, reads as no stamp (not a crash)",
          vx.workbook_stamp(WORK / "missing.xlsx") is None
          and vx.workbook_stamp(DB) is None, "both None")

    print("\n== MyView is yours ==")
    wb = load_workbook(a)
    wb["MyView"]["H20"] = "my own note"
    wb.save(a)
    build(DB, a)
    kept = load_workbook(a)["MyView"]["H20"].value
    check("a rebuild never writes to MyView", kept == "my own note", f"H20 = {kept!r}")

    passed = sum(ok for ok, _ in RESULTS)
    print(f"\n== {passed}/{len(RESULTS)} checks passed ==")
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
